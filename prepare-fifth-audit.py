import csv
import hashlib
import json
import re
import shutil
from pathlib import Path

from docx import Document


AUDIT_ROOT = Path(r"F:\NPI-4-code\讨论\第五次审核")
REFERENCE_DIR = AUDIT_ROOT / "参考文献"
RESULT_DIR = AUDIT_ROOT / "审核结果"
SOURCE_DISCUSSION_PATH = AUDIT_ROOT / "ASD_EC-G1_Discussion_第四次审核_REF顺序更新.docx"
REVISED_DISCUSSION_PATH = RESULT_DIR / "ASD_EC-G1_Discussion_第五次审核_删除错误引用并更新REF.docx"
SOURCE_MAPPING_PATH = AUDIT_ROOT / "第四次引用编号映射_33篇.csv"
FINAL_MAPPING_PATH = RESULT_DIR / "第五次讨论文献最终对应表.csv"
MANIFEST_PATH = RESULT_DIR / "第五次讨论文献重命名清单.json"
STAGE_DIR = REFERENCE_DIR / ".rename-stage-fifth-audit"
ARCHIVE_DIR = REFERENCE_DIR / "archive"

ARCHIVE_FILE_NAMES = {
    "REF9.pdf": "content-mismatch-former-ref9.pdf",
    "REF7.pdf": "not-cited-former-ref7.pdf",
    "nihms-1677334.pdf": "duplicate-of-ref13.pdf",
}


def calculateSha256(filePath: Path) -> str:
    digest = hashlib.sha256()
    with filePath.open("rb") as binaryFile:
        for chunk in iter(lambda: binaryFile.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def readCsvRows(csvPath: Path) -> list[dict[str, str]]:
    with csvPath.open("r", encoding="utf-8-sig", newline="") as csvFile:
        return list(csv.DictReader(csvFile))


def normalizeSourceFileName(rawSourceFile: str) -> str:
    return rawSourceFile.split("（", 1)[0].strip()


def buildFinalMappingRows(sourceRows: list[dict[str, str]]) -> list[dict[str, str]]:
    finalRows: list[dict[str, str]] = []
    for sourceRow in sourceRows:
        fourthRefNumber = int(sourceRow["newRef"].replace("REF", ""))
        if fourthRefNumber == 17:
            continue
        finalRefNumber = fourthRefNumber if fourthRefNumber < 17 else fourthRefNumber - 1
        originalFileName = normalizeSourceFileName(sourceRow["sourceFile"])
        finalRows.append(
            {
                "recordType": "final_reference",
                "finalRef": f"REF{finalRefNumber}",
                "fourthAuditRef": sourceRow["newRef"],
                "previousRef": sourceRow["previousRef"],
                "originalFileName": originalFileName,
                "finalFileName": f"REF{finalRefNumber}.pdf",
                "archiveFileName": "",
                "paragraphs": sourceRow["paragraphs"],
                "ratingSummary": sourceRow["rating"],
                "doi": sourceRow["doi"],
                "year": sourceRow["year"],
                "journal": sourceRow["journal"],
                "title": sourceRow["title"],
                "verificationStatus": "verified_local_full_text",
                "notes": "本地全文题名与DOI一致；已纳入最终连续编号。",
            }
        )
    return finalRows


def reviseDiscussionDocument() -> dict[str, object]:
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(SOURCE_DISCUSSION_PATH, REVISED_DISCUSSION_PATH)
    document = Document(REVISED_DISCUSSION_PATH)
    changedParagraphs: list[dict[str, str]] = []

    for paragraphIndex, paragraph in enumerate(document.paragraphs, start=1):
        originalText = paragraph.text
        revisedText = originalText
        revisedText = re.sub(r",\s*REF17(?=\s*[\]\)])", "", revisedText)
        revisedText = re.sub(r"(?<=[\[\(])REF17,\s*", "", revisedText)
        revisedText = revisedText.replace("[REF17]", "").replace("(REF17)", "")
        revisedText = re.sub(
            r"\bREF(1[89]|2\d|3[0-3])\b",
            lambda match: f"REF{int(match.group(1)) - 1}",
            revisedText,
        )

        if revisedText == originalText:
            continue
        if paragraph.runs:
            paragraph.runs[0].text = revisedText
            for extraRun in paragraph.runs[1:]:
                extraRun.text = ""
        else:
            paragraph.add_run(revisedText)
        changedParagraphs.append(
            {
                "paragraph": f"P{paragraphIndex}",
                "before": originalText,
                "after": revisedText,
            }
        )

    document.core_properties.title = "ASD EC-G1 Discussion — Fifth citation audit revision"
    document.core_properties.subject = "Removed one mismatched citation and compressed REF numbering"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.save(REVISED_DISCUSSION_PATH)

    revisedDocument = Document(REVISED_DISCUSSION_PATH)
    allBodyText = "\n".join(paragraph.text for paragraph in revisedDocument.paragraphs)
    if "(REF16)" not in allBodyText:
        raise RuntimeError("Expected revised P20 citation (REF16) was not found.")
    if "REF33" in allBodyText:
        raise RuntimeError("REF33 remained after numbering compression.")
    citationPlacements = re.findall(r"\bREF\d+\b", allBodyText)
    if len(citationPlacements) != 41:
        raise RuntimeError(f"Expected 41 citation placements, found {len(citationPlacements)}.")
    if max(int(item.replace("REF", "")) for item in citationPlacements) != 32:
        raise RuntimeError("Expected final maximum reference label REF32.")

    return {
        "path": str(REVISED_DISCUSSION_PATH),
        "changedParagraphs": changedParagraphs,
        "citationPlacementCount": len(citationPlacements),
        "uniqueReferenceCount": len(set(citationPlacements)),
    }


def renameReferenceFiles(finalRows: list[dict[str, str]]) -> list[dict[str, str]]:
    canonicalSourceNames = [row["originalFileName"] for row in finalRows]
    if len(canonicalSourceNames) != len(set(canonicalSourceNames)):
        raise RuntimeError("Canonical source filenames are not unique.")

    expectedCurrentNames = set(canonicalSourceNames) | set(ARCHIVE_FILE_NAMES)
    actualCurrentNames = {filePath.name for filePath in REFERENCE_DIR.glob("*.pdf")}
    if actualCurrentNames != expectedCurrentNames:
        missingNames = sorted(expectedCurrentNames - actualCurrentNames)
        unexpectedNames = sorted(actualCurrentNames - expectedCurrentNames)
        raise RuntimeError(
            f"Reference-folder preflight failed. Missing={missingNames}; unexpected={unexpectedNames}"
        )

    if STAGE_DIR.exists():
        raise RuntimeError(f"Staging directory already exists: {STAGE_DIR}")
    STAGE_DIR.mkdir()
    ARCHIVE_DIR.mkdir(exist_ok=True)

    for archiveFileName in ARCHIVE_FILE_NAMES.values():
        if (ARCHIVE_DIR / archiveFileName).exists():
            raise RuntimeError(f"Archive destination already exists: {archiveFileName}")

    renameLog: list[dict[str, str]] = []
    for row in finalRows:
        sourcePath = REFERENCE_DIR / row["originalFileName"]
        stagedPath = STAGE_DIR / row["finalFileName"]
        row["sha256"] = calculateSha256(sourcePath)
        sourcePath.rename(stagedPath)
        renameLog.append(
            {
                "action": "stage_canonical_reference",
                "source": str(sourcePath),
                "destination": str(stagedPath),
            }
        )

    for originalFileName, archiveFileName in ARCHIVE_FILE_NAMES.items():
        sourcePath = REFERENCE_DIR / originalFileName
        archivePath = ARCHIVE_DIR / archiveFileName
        sourcePath.rename(archivePath)
        renameLog.append(
            {
                "action": "archive_nonfinal_reference",
                "source": str(sourcePath),
                "destination": str(archivePath),
                "sha256": calculateSha256(archivePath),
            }
        )

    for row in finalRows:
        stagedPath = STAGE_DIR / row["finalFileName"]
        finalPath = REFERENCE_DIR / row["finalFileName"]
        stagedPath.rename(finalPath)
        renameLog.append(
            {
                "action": "finalize_canonical_reference",
                "source": str(stagedPath),
                "destination": str(finalPath),
            }
        )

    STAGE_DIR.rmdir()
    finalNames = {filePath.name for filePath in REFERENCE_DIR.glob("*.pdf")}
    expectedFinalNames = {f"REF{refNumber}.pdf" for refNumber in range(1, 33)}
    if finalNames != expectedFinalNames:
        raise RuntimeError(f"Final filenames are not REF1–REF32: {sorted(finalNames)}")

    return renameLog


def writeFinalMappingCsv(finalRows: list[dict[str, str]]) -> None:
    outputRows = [dict(row) for row in finalRows]
    outputRows.extend(
        [
            {
                "recordType": "removed_content_mismatch",
                "finalRef": "",
                "fourthAuditRef": "REF17",
                "previousRef": "REF9",
                "originalFileName": "REF9.pdf",
                "finalFileName": "",
                "archiveFileName": "archive/content-mismatch-former-ref9.pdf",
                "paragraphs": "P20（引用已删除）",
                "ratingSummary": "不进入最终评级",
                "doi": "10.2307/26459823（文件实际DOI）；正文原拟引10.1016/j.neuroimage.2021.118546",
                "year": "2005",
                "journal": "Revista de Letras",
                "title": "NOMES DE LUGAR: CONFIM",
                "verificationStatus": "content_mismatch_archived",
                "notes": "题名和DOI均与正文原拟引文献不符；依用户指示删除正文引用。",
                "sha256": "",
            },
            {
                "recordType": "archived_not_cited",
                "finalRef": "",
                "fourthAuditRef": "",
                "previousRef": "REF7",
                "originalFileName": "REF7.pdf",
                "finalFileName": "",
                "archiveFileName": "archive/not-cited-former-ref7.pdf",
                "paragraphs": "",
                "ratingSummary": "不进入最终评级",
                "doi": "10.1038/s41467-025-58187-6",
                "year": "2025",
                "journal": "Nature Communications",
                "title": "Stimulation mapping and whole-brain modeling reveal gradients of excitability and recurrence in cortical networks",
                "verificationStatus": "not_cited_archived",
                "notes": "本地全文正确，但未被最终正文引用。",
                "sha256": "",
            },
            {
                "recordType": "archived_duplicate",
                "finalRef": "",
                "fourthAuditRef": "",
                "previousRef": "nihms-1677334",
                "originalFileName": "nihms-1677334.pdf",
                "finalFileName": "",
                "archiveFileName": "archive/duplicate-of-ref13.pdf",
                "paragraphs": "",
                "ratingSummary": "不进入最终评级",
                "doi": "10.1038/s41583-020-00420-w",
                "year": "2021",
                "journal": "Nature Reviews Neuroscience",
                "title": "The default mode network: where the idiosyncratic self meets the shared social world",
                "verificationStatus": "duplicate_full_text_archived",
                "notes": "与最终REF13为同一篇文献的另一版本。",
                "sha256": "",
            },
        ]
    )

    fieldNames = [
        "recordType",
        "finalRef",
        "fourthAuditRef",
        "previousRef",
        "originalFileName",
        "finalFileName",
        "archiveFileName",
        "paragraphs",
        "ratingSummary",
        "doi",
        "year",
        "journal",
        "title",
        "verificationStatus",
        "sha256",
        "notes",
    ]
    with FINAL_MAPPING_PATH.open("w", encoding="utf-8-sig", newline="") as csvFile:
        writer = csv.DictWriter(csvFile, fieldnames=fieldNames)
        writer.writeheader()
        writer.writerows(outputRows)


def main() -> None:
    sourceRows = readCsvRows(SOURCE_MAPPING_PATH)
    finalRows = buildFinalMappingRows(sourceRows)
    if len(finalRows) != 32:
        raise RuntimeError(f"Expected 32 final references, found {len(finalRows)}.")

    discussionRevision = reviseDiscussionDocument()
    renameLog = renameReferenceFiles(finalRows)
    writeFinalMappingCsv(finalRows)

    manifest = {
        "discussionRevision": discussionRevision,
        "finalReferenceCount": len(finalRows),
        "archivedReferenceCount": len(ARCHIVE_FILE_NAMES),
        "finalMappingPath": str(FINAL_MAPPING_PATH),
        "renameLog": renameLog,
    }
    MANIFEST_PATH.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
